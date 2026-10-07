"""Slides previewed through LibreOffice as a PDF (with the HTML fallback).

The real `soffice` is not part of the test environment, so a tiny shell
script stands in for it: it honours `--outdir` and writes `<stem>.pdf`,
counting its invocations so the content-hash cache can be asserted.
"""
import os
import stat
from pathlib import Path

import pytest

from app.core.config import Settings
from app.domains.workspace.office_pdf import OfficeConversionError, OfficePdfConverter
from app.domains.workspace.service import WorkspaceService

FAKE_SOFFICE = """#!/bin/sh
# minimal stand-in for `soffice --headless --convert-to pdf --outdir D F`
outdir=""
input=""
while [ $# -gt 0 ]; do
  case "$1" in
    --outdir) outdir="$2"; shift 2 ;;
    --convert-to) shift 2 ;;
    -*) shift ;;
    *) input="$1"; shift ;;
  esac
done
echo run >> "$(dirname "$0")/calls.log"
if [ -n "$FAKE_SOFFICE_FAIL" ]; then echo "boom" >&2; exit 1; fi
stem=$(basename "$input")
stem="${stem%.*}"
printf '%%PDF-1.4 fake from %s\\n' "$input" > "$outdir/$stem.pdf"
"""


def _fake_soffice(tmp_path: Path) -> Path:
    exe = tmp_path / "bin" / "soffice"
    exe.parent.mkdir(parents=True, exist_ok=True)
    exe.write_text(FAKE_SOFFICE)
    exe.chmod(exe.stat().st_mode | stat.S_IEXEC)
    return exe


def _calls(exe: Path) -> int:
    log = exe.parent / "calls.log"
    return len(log.read_text().splitlines()) if log.exists() else 0


def _deck(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Hello"
    prs.save(str(path))


def _settings(root: Path, converter: str) -> Settings:
    return Settings(
        content_root=root,
        data_dir=root.parent / ".noteeli",
        session_secret="test-secret",
        google_client_id="",
        google_client_secret="",
        office_converter=converter,
    )


def test_converter_runs_libreoffice_and_caches_by_content(tmp_path: Path):
    exe = _fake_soffice(tmp_path)
    converter = OfficePdfConverter(_settings(tmp_path / "vault", str(exe)))
    assert converter.available()

    first = converter.convert_to_pdf(b"deck-one", ".pptx")
    assert first.is_file() and first.suffix == ".pdf"
    assert first.read_bytes().startswith(b"%PDF")
    assert first.parent == tmp_path / ".noteeli" / "office-pdf-cache"

    again = converter.convert_to_pdf(b"deck-one", ".pptx")
    assert again == first
    assert _calls(exe) == 1, "the same bytes must not be converted twice"

    other = converter.convert_to_pdf(b"deck-two", ".pptx")
    assert other != first
    assert _calls(exe) == 2


def test_converter_absent_or_off(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("PATH", str(tmp_path / "empty"))
    assert OfficePdfConverter(_settings(tmp_path / "vault", "")).available() is False
    assert OfficePdfConverter(_settings(tmp_path / "vault", "off")).available() is False
    with pytest.raises(OfficeConversionError):
        OfficePdfConverter(_settings(tmp_path / "vault", "off")).convert_to_pdf(b"x", ".pptx")


def test_converter_failure_raises_not_caches(tmp_path: Path, monkeypatch):
    exe = _fake_soffice(tmp_path)
    monkeypatch.setenv("FAKE_SOFFICE_FAIL", "1")
    converter = OfficePdfConverter(_settings(tmp_path / "vault", str(exe)))
    with pytest.raises(OfficeConversionError):
        converter.convert_to_pdf(b"deck", ".pptx")
    assert not list(converter.cache_dir.glob("*.pdf"))


def test_read_document_reports_pdf_rendering_only_with_a_converter(tmp_path: Path):
    notes = tmp_path / "vault"
    notes.mkdir()
    _deck(notes / "deck.pptx")
    exe = _fake_soffice(tmp_path)

    with_converter = WorkspaceService(_settings(notes, str(exe))).read_document("deck.pptx")
    assert with_converter.preview_kind == "pptx"
    assert with_converter.preview_rendering == "pdf"

    without = WorkspaceService(_settings(notes, "off")).read_document("deck.pptx")
    assert without.preview_rendering == "html"

    # Word stays an HTML preview even when LibreOffice exists.
    (notes / "x.docx").write_bytes(b"PK")
    assert WorkspaceService(_settings(notes, str(exe))).get_preview_rendering("docx") == "html"
    assert WorkspaceService(_settings(notes, str(exe))).get_preview_rendering("image") is None


def test_render_office_pdf_goes_through_the_storage_backend(tmp_path: Path):
    notes = tmp_path / "vault"
    notes.mkdir()
    _deck(notes / "deck.pptx")
    exe = _fake_soffice(tmp_path)
    service = WorkspaceService(_settings(notes, str(exe)))
    pdf = service.render_office_pdf("deck.pptx")
    assert pdf.read_bytes().startswith(b"%PDF")
    with pytest.raises(Exception):
        service.render_office_pdf("../deck.pptx")


# ── the preview endpoint ──────────────────────────────────────────────
# The workspace router snapshots its service at import time against the
# conftest sandbox, so the converter is swapped on that singleton and the
# deck is written into the sandbox content root.


@pytest.fixture
def preview_client(tmp_path: Path, monkeypatch):
    from fastapi.testclient import TestClient

    from app.domains.workspace import router as workspace_router
    from app.main import create_app

    content = Path(os.environ["NOTEELI_CONTENT_ROOT"])
    content.mkdir(parents=True, exist_ok=True)
    _deck(content / "office-pdf-deck.pptx")
    exe = _fake_soffice(tmp_path)
    original = workspace_router.workspace_service.office_pdf
    converter = OfficePdfConverter(_settings(content, str(exe)))
    # The cache is keyed by content: a PDF left by an earlier test would
    # short-circuit the conversion this test wants to observe.
    for stale in converter.cache_dir.glob("*.pdf"):
        stale.unlink()
    workspace_router.workspace_service.office_pdf = converter
    try:
        yield TestClient(create_app(), base_url="http://127.0.0.1"), exe
    finally:
        workspace_router.workspace_service.office_pdf = original
        (content / "office-pdf-deck.pptx").unlink(missing_ok=True)


def test_preview_endpoint_serves_slides_as_pdf(preview_client):
    client, _exe = preview_client
    doc = client.get("/api/file", params={"path": "office-pdf-deck.pptx"}).json()
    assert doc["preview_kind"] == "pptx" and doc["preview_rendering"] == "pdf"

    response = client.get("/api/file/preview", params={"path": "office-pdf-deck.pptx"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF")


def test_preview_endpoint_falls_back_to_html_when_libreoffice_fails(preview_client, monkeypatch):
    client, _exe = preview_client
    monkeypatch.setenv("FAKE_SOFFICE_FAIL", "1")
    response = client.get("/api/file/preview", params={"path": "office-pdf-deck.pptx"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "Slajd 1" in response.text
