"""E2E (headless Chromium): a .pptx opened without LibreOffice shows the
fallback banner with the in-app download offer instead of silently
rendering text cards.

Same caveats as the other e2e modules (CDN + Playwright Chromium).
"""
import platform
import sys
from pathlib import Path

import pytest

playwright_api = pytest.importorskip("playwright.sync_api")
from playwright.sync_api import expect, sync_playwright  # noqa: E402

from test_e2e_git_badge import _cdn_reachable, _running_server  # noqa: E402


def _deck(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Banner deck"
    prs.save(str(path))


def test_pptx_without_converter_shows_download_banner(tmp_path):
    if not _cdn_reachable():
        pytest.skip("editor CDN unreachable — cannot load the real UI")
    with _running_server(tmp_path, extra_env={"NOTEELI_OFFICE_CONVERTER": "off"}) as (base_url, content):
        _deck(content / "deck.pptx")
        with sync_playwright() as p:
            try:
                browser = p.chromium.launch(headless=True)
            except Exception as exc:
                pytest.skip(f"Playwright Chromium unavailable: {exc}")
            try:
                page = browser.new_page()
                page.goto(base_url, wait_until="networkidle")
                page.locator(".tree-link-file", has_text="deck.pptx").first.click()
                banner = page.locator("#office-fallback-banner")
                expect(banner).to_be_visible(timeout=15_000)
                expect(page.locator("#office-fallback-text")).to_contain_text("LibreOffice")
                portable = sys.platform.startswith("linux") and platform.machine() in ("x86_64", "AMD64")
                if portable:
                    expect(page.locator("#office-fallback-install")).to_be_visible()
                else:
                    expect(page.locator("#office-fallback-command")).to_contain_text("libreoffice")
                # The text-card fallback is still shown underneath.
                expect(page.locator("#office-preview")).to_be_visible()
            finally:
                browser.close()
